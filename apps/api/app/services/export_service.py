import html
import re
from datetime import datetime, timezone
from typing import Literal, Optional

from docx import Document
from pydantic import BaseModel

from app.repositories import export_repo
from app.services import resolved_template_service


class ExportDocument(BaseModel):
    ok: bool = True
    task_id: str
    format: Literal["markdown", "html", "word"]
    generated_at: str
    file_path: str
    filename: str
    content: Optional[str] = None


def _slugify(value: str) -> str:
    slug = re.sub(r"[^\w\u4e00-\u9fff\- ]+", "", value).strip().lower()
    slug = slug.replace(" ", "-")
    return slug or "section"


def _sort_resolved_nodes(
    nodes: list[resolved_template_service.ResolvedNode],
) -> list[resolved_template_service.ResolvedNode]:
    if not nodes:
        return []
    index_by_id = {node.node_id: idx for idx, node in enumerate(nodes)}
    children_by_parent: dict[Optional[str], list[resolved_template_service.ResolvedNode]] = {}
    for node in nodes:
        parent = node.parent_id if node.parent_id in index_by_id else None
        children_by_parent.setdefault(parent, []).append(node)
    for parent_id, items in children_by_parent.items():
        # Keep section tree order stable, but place same-title fill fields
        # immediately below their corresponding section title.
        section_order_by_title: dict[str, int] = {}
        for item in items:
            if item.node_type == "section":
                section_order_by_title.setdefault(
                    item.node_title,
                    index_by_id.get(item.node_id, 0),
                )

        def child_sort_key(node: resolved_template_service.ResolvedNode) -> tuple[int, int, int]:
            original = index_by_id.get(node.node_id, 0)
            if node.node_type == "section":
                return (original, 0, original)
            if (
                node.node_type == "field"
                and node.content_mode == "fill"
                and node.node_title in section_order_by_title
                and parent_id is not None
            ):
                # same-title field under same parent should render right below section title
                return (section_order_by_title[node.node_title], 1, original)
            return (original, 2, original)

        items.sort(key=child_sort_key)

    ordered: list[resolved_template_service.ResolvedNode] = []

    def dfs(parent_id: Optional[str]) -> None:
        for child in children_by_parent.get(parent_id, []):
            ordered.append(child)
            dfs(child.node_id)

    dfs(None)
    seen = {node.node_id for node in ordered}
    for node in nodes:
        if node.node_id not in seen:
            ordered.append(node)
    return ordered


def _render_node_markdown(node: resolved_template_service.ResolvedNode) -> list[str]:
    payload = node.render_payload or {}
    payload_type = str(payload.get("type") or "")
    lines: list[str] = []

    if node.node_type == "section":
        heading_level = min(max(node.level + 1, 2), 6)
        lines.append(f'{"#" * heading_level} {node.node_title}')

        if payload_type == "rich_section":
            text = str(payload.get("text") or "").strip()
            if text:
                lines.append(text)
            images = payload.get("images") or []
            if isinstance(images, list):
                for image in images:
                    if not isinstance(image, dict):
                        continue
                    caption = str(image.get("caption") or "section-image").strip() or "section-image"
                    path = str(image.get("path") or "").strip()
                    unit_id = str(image.get("unit_id") or "").strip()
                    if path:
                        lines.append(f"![{caption}]({path})")
                    else:
                        lines.append(f"![{caption}](#image:{unit_id})")
        elif payload_type == "mixed_section":
            text = str(payload.get("text") or "").strip()
            if text:
                lines.append(text)
            fields = payload.get("fields") or []
            if isinstance(fields, list) and fields:
                lines.append("")
                lines.append("**字段回填结果**")
                for field in fields:
                    if not isinstance(field, dict):
                        continue
                    field_title = str(field.get("node_title") or "field")
                    field_status = str(field.get("fill_status") or "unknown")
                    field_payload = field.get("render_payload")
                    lines.append(
                        f"- {field_title} ({field_status}): {field_payload if field_payload is not None else '(empty)'}"
                    )
            images = payload.get("images") or []
            if isinstance(images, list) and images:
                lines.append("")
                lines.append("**章节附图**")
                for image in images:
                    if not isinstance(image, dict):
                        continue
                    caption = str(image.get("caption") or "section-image").strip() or "section-image"
                    path = str(image.get("path") or "").strip()
                    unit_id = str(image.get("unit_id") or "").strip()
                    if path:
                        lines.append(f"![{caption}]({path})")
                    else:
                        lines.append(f"![{caption}](#image:{unit_id})")
        elif payload_type == "text":
            value = str(payload.get("value") or "").strip()
            if value:
                lines.append(value)
    else:
        if payload_type == "image":
            caption = str(payload.get("title") or node.node_title).strip() or node.node_title
            path = str(payload.get("path") or "").strip()
            unit_id = str(payload.get("unit_id") or "").strip()
            if path:
                lines.append(f"- **{node.node_title}**")
                lines.append(f"  ![{caption}]({path})")
            else:
                lines.append(f"- **{node.node_title}**: [image:{unit_id}]")
        elif payload_type in {"text", "list"}:
            value = str(payload.get("value") or "").strip()
            lines.append(f"- **{node.node_title}**: {value or '(empty)'}")
        elif payload_type == "table_ref":
            lines.append(f"- **{node.node_title}**: `{payload}`")
        else:
            lines.append(f"- **{node.node_title}**: `(unresolved)`")

    if node.resolved_status in {"risky", "missing"}:
        lines.append(f"> 状态: {node.resolved_status}，建议人工复核。")
    if node.gap_count > 0:
        lines.append(f"> 关联缺口数: {node.gap_count}")
    return lines


def _build_mirror_fill_by_section_id(
    nodes: list[resolved_template_service.ResolvedNode],
) -> tuple[dict[str, resolved_template_service.ResolvedNode], set[str]]:
    section_by_parent_and_title: dict[tuple[Optional[str], str], resolved_template_service.ResolvedNode] = {}
    for node in nodes:
        if node.node_type != "section":
            continue
        section_by_parent_and_title[(node.parent_id, node.node_title)] = node

    mirror_by_section_id: dict[str, resolved_template_service.ResolvedNode] = {}
    consumed_field_ids: set[str] = set()
    for node in nodes:
        if node.node_type != "field":
            continue
        section = section_by_parent_and_title.get((node.parent_id, node.node_title))
        if not section:
            continue
        if section.node_id in mirror_by_section_id:
            continue
        mirror_by_section_id[section.node_id] = node
        consumed_field_ids.add(node.node_id)
    return mirror_by_section_id, consumed_field_ids


def _render_field_payload_inline(
    field_node: resolved_template_service.ResolvedNode,
) -> list[str]:
    payload = field_node.render_payload or {}
    payload_type = str(payload.get("type") or "")
    lines: list[str] = []
    if payload_type == "image":
        caption = str(payload.get("title") or field_node.node_title).strip() or field_node.node_title
        path = str(payload.get("path") or "").strip()
        unit_id = str(payload.get("unit_id") or "").strip()
        if path:
            lines.append(f"![{caption}]({path})")
        else:
            lines.append(f"[image:{unit_id}]")
    elif payload_type in {"text", "list"}:
        value = str(payload.get("value") or "").strip()
        lines.append(value or "unresolved")
    elif payload_type == "table_ref":
        lines.append(f"`{payload}`")
    else:
        lines.append("unresolved")
    if field_node.resolved_status in {"risky", "missing"}:
        lines.append(f"> 状态: {field_node.resolved_status}，建议人工复核。")
    return lines


def render_resolved_template_markdown(
    resolved: resolved_template_service.ResolvedTemplate,
) -> str:
    nodes = _sort_resolved_nodes(resolved.nodes)
    mirror_by_section_id, consumed_field_ids = _build_mirror_fill_by_section_id(nodes)
    lines: list[str] = [
        f"# 项目交付草案（任务 {resolved.task_id}）",
        "",
        "## 结果摘要",
        f"- total_nodes: {resolved.summary.total_nodes}",
        f"- ready_nodes: {resolved.summary.ready_nodes}",
        f"- risky_nodes: {resolved.summary.risky_nodes}",
        f"- missing_nodes: {resolved.summary.missing_nodes}",
        f"- export_readiness: {resolved.summary.export_readiness}",
        "",
    ]

    section_nodes = [node for node in nodes if node.node_type == "section"]
    if section_nodes:
        lines.append("## 目录")
        lines.append("")
        for section in section_nodes:
            indent = "  " * max(0, int(section.level) - 1)
            anchor = _slugify(section.node_title)
            lines.append(f"{indent}- [{section.node_title}](#{anchor})")
        lines.append("")
        lines.append("## 正文")
        lines.append("")

    appendix_fields: list[resolved_template_service.ResolvedNode] = []
    for node in nodes:
        if node.node_type == "field" and node.node_id in consumed_field_ids:
            continue
        node_lines = _render_node_markdown(node)
        if node.node_type == "section":
            mirror_field = mirror_by_section_id.get(node.node_id)
            if mirror_field:
                inline_lines = _render_field_payload_inline(mirror_field)
                if inline_lines:
                    node_lines.extend(inline_lines)
        elif node.node_type == "field":
            appendix_fields.append(node)
            continue
        if node_lines:
            lines.extend(node_lines)
            lines.append("")

    if appendix_fields:
        lines.append("## 字段回填附录")
        lines.append("")
        for field_node in appendix_fields:
            lines.append(f"- {field_node.node_title}")
            field_lines = _render_field_payload_inline(field_node)
            for field_line in field_lines:
                lines.append(f"  {field_line}")
        lines.append("")

    if resolved.gaps:
        lines.append("## 缺口与风险清单")
        lines.append("")
        for gap in resolved.gaps:
            lines.append(
                f"- [{gap.severity}] {gap.node_title}: {gap.description} "
                f"(action: {gap.manual_action or 'n/a'})"
            )
        lines.append("")

    return "\n".join(lines).strip() + "\n"


def _markdown_to_html(markdown: str, title: str) -> str:
    lines = markdown.splitlines()
    html_lines: list[str] = [
        "<!doctype html>",
        "<html><head><meta charset='utf-8'/>",
        f"<title>{html.escape(title)}</title>",
        "<style>body{font-family:Arial,sans-serif;max-width:980px;margin:24px auto;line-height:1.6;padding:0 12px} pre{white-space:pre-wrap}</style>",
        "</head><body>",
    ]

    in_list = False
    for raw in lines:
        line = raw.rstrip()
        if not line:
            if in_list:
                html_lines.append("</ul>")
                in_list = False
            continue
        if line.startswith("#"):
            if in_list:
                html_lines.append("</ul>")
                in_list = False
            level = min(len(line) - len(line.lstrip("#")), 6)
            text = line[level:].strip()
            html_lines.append(f"<h{level}>{html.escape(text)}</h{level}>")
            continue
        if line.startswith("- "):
            if not in_list:
                html_lines.append("<ul>")
                in_list = True
            html_lines.append(f"<li>{html.escape(line[2:].strip())}</li>")
            continue
        image_match = re.match(r"!\[(?P<alt>.*?)\]\((?P<src>.*?)\)", line)
        if image_match:
            if in_list:
                html_lines.append("</ul>")
                in_list = False
            alt = html.escape(image_match.group("alt"))
            src = html.escape(image_match.group("src"))
            html_lines.append(f"<p><img src='{src}' alt='{alt}' style='max-width:100%'/></p>")
            continue
        if line.startswith("> "):
            if in_list:
                html_lines.append("</ul>")
                in_list = False
            html_lines.append(f"<blockquote>{html.escape(line[2:])}</blockquote>")
            continue
        if in_list:
            html_lines.append("</ul>")
            in_list = False
        html_lines.append(f"<p>{html.escape(line)}</p>")

    if in_list:
        html_lines.append("</ul>")
    html_lines.append("</body></html>")
    return "\n".join(html_lines)


def export_markdown(task_id: str, rebuild: bool = True) -> ExportDocument:
    resolved = (
        resolved_template_service.build_resolved_template(task_id)
        if rebuild
        else resolved_template_service.get_resolved_template(task_id)
    )
    markdown = render_resolved_template_markdown(resolved)
    path = export_repo.save_export_content(task_id, "markdown", markdown)
    generated_at = datetime.now(timezone.utc).isoformat()
    return ExportDocument(
        ok=True,
        task_id=task_id,
        format="markdown",
        generated_at=generated_at,
        file_path=str(path),
        filename=path.name,
        content=markdown,
    )


def get_markdown_export(task_id: str) -> ExportDocument:
    content = export_repo.load_export_content(task_id, "markdown")
    if content is None:
        raise FileNotFoundError(f"markdown export not found for task {task_id}")
    path = export_repo.export_path(task_id, "markdown")
    return ExportDocument(
        ok=True,
        task_id=task_id,
        format="markdown",
        generated_at=datetime.now(timezone.utc).isoformat(),
        file_path=str(path),
        filename=path.name,
        content=content,
    )


def export_html(task_id: str, rebuild: bool = False) -> ExportDocument:
    markdown_doc = export_markdown(task_id, rebuild=rebuild)
    html_content = _markdown_to_html(markdown_doc.content, f"Task {task_id} Export")
    path = export_repo.save_export_content(task_id, "html", html_content)
    return ExportDocument(
        ok=True,
        task_id=task_id,
        format="html",
        generated_at=datetime.now(timezone.utc).isoformat(),
        file_path=str(path),
        filename=path.name,
        content=html_content,
    )


def get_html_export(task_id: str) -> ExportDocument:
    content = export_repo.load_export_content(task_id, "html")
    if content is None:
        raise FileNotFoundError(f"html export not found for task {task_id}")
    path = export_repo.export_path(task_id, "html")
    return ExportDocument(
        ok=True,
        task_id=task_id,
        format="html",
        generated_at=datetime.now(timezone.utc).isoformat(),
        file_path=str(path),
        filename=path.name,
        content=content,
    )


def _apply_node_to_doc(
    document: Document,
    node: resolved_template_service.ResolvedNode,
    mirror_field: Optional[resolved_template_service.ResolvedNode] = None,
) -> None:
    payload = node.render_payload or {}
    payload_type = str(payload.get("type") or "")

    if node.node_type == "section":
        heading_level = min(max(int(node.level), 1), 9)
        document.add_heading(node.node_title, level=heading_level)
        if payload_type in {"rich_section", "mixed_section"}:
            text = str(payload.get("text") or "").strip()
            if text:
                document.add_paragraph(text)
            images = payload.get("images") or []
            if isinstance(images, list):
                for image in images:
                    if not isinstance(image, dict):
                        continue
                    caption = str(image.get("caption") or "").strip()
                    path = str(image.get("path") or "").strip()
                    if path:
                        document.add_paragraph(f"[Image] {path}")
                    if caption:
                        document.add_paragraph(f"Caption: {caption}")
            if payload_type == "mixed_section":
                fields = payload.get("fields") or []
                if isinstance(fields, list) and fields:
                    document.add_paragraph("字段回填结果")
                    for field in fields:
                        if not isinstance(field, dict):
                            continue
                        title = str(field.get("node_title") or "field")
                        status = str(field.get("fill_status") or "unknown")
                        document.add_paragraph(f"- {title} ({status})")
        elif payload_type == "text":
            value = str(payload.get("value") or "").strip()
            if value:
                document.add_paragraph(value)
        if mirror_field is not None:
            mirror_payload = mirror_field.render_payload or {}
            mirror_type = str(mirror_payload.get("type") or "")
            if mirror_type == "image":
                path = str(mirror_payload.get("path") or "").strip()
                unit_id = str(mirror_payload.get("unit_id") or "").strip()
                if path:
                    document.add_paragraph(f"[Image] {path}")
                else:
                    document.add_paragraph(f"[Image] image:{unit_id}")
            elif mirror_type in {"text", "list"}:
                value = str(mirror_payload.get("value") or "").strip()
                document.add_paragraph(value or "unresolved")
            elif mirror_type == "table_ref":
                document.add_paragraph(str(mirror_payload))
            else:
                document.add_paragraph("unresolved")
        return

    if payload_type == "image":
        title = str(payload.get("title") or node.node_title).strip() or node.node_title
        path = str(payload.get("path") or "").strip()
        if path:
            document.add_paragraph(f"{title}: {path}")
        else:
            unit_id = str(payload.get("unit_id") or "").strip()
            document.add_paragraph(f"{title}: image:{unit_id}")
        return

    if payload_type == "text":
        value = str(payload.get("value") or "").strip() or "(empty)"
        document.add_paragraph(f"{node.node_title}: {value}")
        return

    document.add_paragraph(f"{node.node_title}: unresolved")


def export_word(task_id: str, rebuild: bool = True) -> ExportDocument:
    resolved = (
        resolved_template_service.build_resolved_template(task_id)
        if rebuild
        else resolved_template_service.get_resolved_template(task_id)
    )
    nodes = _sort_resolved_nodes(resolved.nodes)
    mirror_by_section_id, consumed_field_ids = _build_mirror_fill_by_section_id(nodes)
    document = Document()
    document.add_heading(f"项目交付草案（任务 {resolved.task_id}）", level=0)
    document.add_paragraph(
        f"summary: total={resolved.summary.total_nodes}, ready={resolved.summary.ready_nodes}, "
        f"risky={resolved.summary.risky_nodes}, missing={resolved.summary.missing_nodes}, "
        f"readiness={resolved.summary.export_readiness}"
    )
    appendix_fields: list[resolved_template_service.ResolvedNode] = []
    section_nodes = [node for node in nodes if node.node_type == "section"]
    if section_nodes:
        document.add_heading("目录", level=1)
        for section in section_nodes:
            indent = "  " * max(0, int(section.level) - 1)
            document.add_paragraph(f"{indent}{section.node_title}")
        # Explicitly split TOC page and body page.
        document.add_page_break()
        document.add_heading("正文", level=1)

    for node in nodes:
        if node.node_type == "field" and node.node_id in consumed_field_ids:
            continue
        if node.node_type == "field":
            appendix_fields.append(node)
            continue
        _apply_node_to_doc(document, node, mirror_field=mirror_by_section_id.get(node.node_id))
    if appendix_fields:
        document.add_heading("字段回填附录", level=1)
        for field_node in appendix_fields:
            document.add_paragraph(field_node.node_title)
            payload = field_node.render_payload or {}
            payload_type = str(payload.get("type") or "")
            if payload_type == "image":
                path = str(payload.get("path") or "").strip()
                unit_id = str(payload.get("unit_id") or "").strip()
                document.add_paragraph(f"[Image] {path or f'image:{unit_id}'}")
            elif payload_type in {"text", "list"}:
                document.add_paragraph(str(payload.get("value") or "unresolved"))
            elif payload_type == "table_ref":
                document.add_paragraph(str(payload))
            else:
                document.add_paragraph("unresolved")
    if resolved.gaps:
        document.add_heading("缺口与风险清单", level=1)
        for gap in resolved.gaps:
            document.add_paragraph(
                f"[{gap.severity}] {gap.node_title}: {gap.description} (action: {gap.manual_action or 'n/a'})"
            )

    path = export_repo.export_path(task_id, "word")
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))
    return ExportDocument(
        ok=True,
        task_id=task_id,
        format="word",
        generated_at=datetime.now(timezone.utc).isoformat(),
        file_path=str(path),
        filename=path.name,
        content=None,
    )


def get_word_export(task_id: str) -> ExportDocument:
    path = export_repo.export_path(task_id, "word")
    if not path.exists():
        raise FileNotFoundError(f"word export not found for task {task_id}")
    return ExportDocument(
        ok=True,
        task_id=task_id,
        format="word",
        generated_at=datetime.now(timezone.utc).isoformat(),
        file_path=str(path),
        filename=path.name,
        content=None,
    )
